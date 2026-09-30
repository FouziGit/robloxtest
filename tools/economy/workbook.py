"""Builds docs/economy/Vellum-economie.xlsx from the pacing simulator's export.

    lune run scripts/pacing -- --csv
    python3 tools/economy/workbook.py

The CSV (docs/economy/pacing.csv, one row per profile and day) is copied as it is into the sheet
"Données"; every other number in the workbook is a formula over that sheet, so pasting a fresher export
over "Données" updates the synthesis and the curves without running this again. Nothing here computes a
figure itself: the simulator (scripts/pacing.luau) is the only source, and tests/Pacing.spec.luau holds it.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs" / "economy" / "pacing.csv"
OUTPUT = ROOT / "docs" / "economy" / "Vellum-economie.xlsx"

FONT = "Arial"
INPUT = Font(name=FONT, color="0000FF")
PLAIN = Font(name=FONT, color="000000")
HEADER = Font(name=FONT, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, bold=True, size=14)
NOTE = Font(name=FONT, italic=True, color="555555")
BAND = PatternFill("solid", start_color="1F2A44")

PROFILES = [
    ("Solo20", "20 min/j seul"),
    ("Pvp20", "20 min/j en PvP"),
    ("Pvp60", "60 min/j en PvP"),
    ("Farmer", "Fermier d'Épreuves"),
]
MAX_TIER = 50
LEVELS = [5, 10, 20, 40]


def header(sheet, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, start=1):
        cell = sheet.cell(row=row, column=column, value=label)
        cell.font = HEADER
        cell.fill = BAND
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def data_sheet(book: Workbook) -> tuple[str, int]:
    sheet = book.active
    sheet.title = "Données"
    with SOURCE.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle))
    header(sheet, 1, rows[0])
    for index, row in enumerate(rows[1:], start=2):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=index, column=column, value=value if column == 1 else int(value))
            cell.font = INPUT
    sheet["A1"].comment = Comment(
        "Source : scripts/pacing.luau (lune run scripts/pacing -- --csv), docs/economy/pacing.csv.", "Vellum"
    )
    for column in range(1, len(rows[0]) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 15
    sheet.freeze_panes = "A2"
    return "Données", len(rows)


def synthesis_sheet(book: Workbook, last: int) -> None:
    sheet = book.create_sheet("Synthèse")
    sheet["A1"] = "Rythme de Vellum sur 60 jours (simulateur, configs réelles)"
    sheet["A1"].font = TITLE
    sheet["A2"] = "Chaque chiffre est une formule sur l'onglet Données ; « — » : pas atteint en 60 jours."
    sheet["A2"].font = NOTE
    labels = ["Profil", "Id", "Pass fini (jour)"]
    labels += [f"Niveau {level} (jour)" for level in LEVELS]
    labels += ["Folios gagnés J7", "Folios gagnés J30", "Niveau J60", "Palier J60", "Objets J60", "Minutes J60"]
    header(sheet, 4, labels)
    profile = f"'Données'!$A$2:$A${last}"
    day = f"'Données'!$B$2:$B${last}"

    def column(letter: str) -> str:
        return f"'Données'!${letter}$2:${letter}${last}"

    for row, (identifier, name) in enumerate(PROFILES, start=5):
        sheet.cell(row=row, column=1, value=name).font = PLAIN
        sheet.cell(row=row, column=2, value=identifier).font = INPUT
        key = f"$B{row}"
        # Tiers and levels only climb: the first day at a mark is one more than the days below it.
        below = f'COUNTIFS({profile},{key},{column("F")},"<{MAX_TIER}")'
        total = f"COUNTIFS({profile},{key})"
        sheet.cell(row=row, column=3, value=f'=IF({below}={total},"—",{below}+1)')
        for offset, level in enumerate(LEVELS):
            under = f'COUNTIFS({profile},{key},{column("D")},"<{level}")'
            sheet.cell(row=row, column=4 + offset, value=f'=IF({under}={total},"—",{under}+1)')
        for offset, (letter, target) in enumerate([("G", 7), ("G", 30), ("D", 60), ("F", 60), ("I", 60), ("C", 60)]):
            formula = f"=SUMIFS({column(letter)},{profile},{key},{day},{target})"
            sheet.cell(row=row, column=8 + offset, value=formula)
        for column_index in range(3, len(labels) + 1):
            cell = sheet.cell(row=row, column=column_index)
            cell.font = PLAIN
            cell.number_format = "#,##0;(#,##0);-"
            cell.alignment = Alignment(horizontal="center")
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 9
    for column_index in range(3, len(labels) + 1):
        sheet.column_dimensions[get_column_letter(column_index)].width = 14
    sheet.row_dimensions[4].height = 32
    notes = [
        "Les minutes jusqu'aux niveaux 5/10/20/40, l'XP/h par activité et le détail des hypothèses de jeu sont",
        "dans docs/ECONOMY.md §2 bis (sortie de lune run scripts/pacing -- --table).",
    ]
    for offset, line in enumerate(notes, start=11):
        sheet.cell(row=offset, column=1, value=line).font = NOTE


def curves_sheet(book: Workbook, last: int) -> None:
    sheet = book.create_sheet("Courbes")
    header(sheet, 1, ["Jour"] + [name for _, name in PROFILES])
    profile = f"'Données'!$A$2:$A${last}"
    day = f"'Données'!$B$2:$B${last}"
    tier = f"'Données'!$F$2:$F${last}"
    for row in range(2, 62):
        sheet.cell(row=row, column=1, value=f"=ROW()-1").font = PLAIN
        for column, (identifier, _) in enumerate(PROFILES, start=2):
            formula = f'=SUMIFS({tier},{profile},"{identifier}",{day},$A{row})'
            sheet.cell(row=row, column=column, value=formula).font = PLAIN
    for column in range(1, len(PROFILES) + 2):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    chart = LineChart()
    chart.title = "Palier du pass, jour après jour"
    chart.y_axis.title = "Palier"
    chart.x_axis.title = "Jour"
    chart.height = 9
    chart.width = 18
    chart.add_data(Reference(sheet, min_col=2, max_col=len(PROFILES) + 1, min_row=1, max_row=61), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=61))
    sheet.add_chart(chart, "G2")


def main() -> None:
    book = Workbook()
    _, last = data_sheet(book)
    synthesis_sheet(book, last)
    curves_sheet(book, last)
    book.active = book.sheetnames.index("Synthèse")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    book.save(OUTPUT)
    print(f"wrote {OUTPUT.relative_to(ROOT)} ({last - 1} rows)")


if __name__ == "__main__":
    main()
